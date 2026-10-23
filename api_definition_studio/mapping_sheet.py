"""Mapping sheet: export the model to Excel (and CSV), and import it back.

The workbook is a round trip: edit it in Excel, import it, and the model,
OpenAPI and diagram all follow. Grey columns are generated and ignored on
import; everything else is read back by its header, so columns can be
reordered."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__, naming
from .model import ATTRIBUTE_TYPES, CARDINALITIES, EXPOSE, OPERATIONS, SECURITY, clean_tags, derived_from
from .source_kinds import KINDS
from .openapi_gen import endpoints, navs_with_target
from .store import Invalid, new_id

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
CALC_FILL = PatternFill("solid", fgColor="5B6B7F")
CALC_CELL = PatternFill("solid", fgColor="EEF1F5")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (header, key, width, generated?)
NS_COLS = [("Key", "key", 10, False), ("Title", "title", 28, False), ("Base path", "base_path", 14, False),
           ("Base path placement", "base_path_in", 18, False), ("Version", "version", 10, False),
           ("Security", "security", 12, False), ("Owner", "owner", 20, False),
           ("Servers (url | description, one per line)", "servers", 48, False),
           ("Description", "description", 50, False), ("Enum code style", "enum_style", 16, False),
           ("Source mapping in OpenAPI", "source_in_openapi", 16, False)]
SS_COLS = [("Key", "key", 12, False), ("Name", "name", 24, False), ("Type", "kind", 12, False),
           ("Owner", "owner", 20, False), ("Internal name tokens (comma separated)", "internal_tokens", 36, False),
           ("Description", "description", 50, False)]
EN_COLS = [("Namespace", "namespace", 12, False), ("Source system", "source_system", 14, False),
           ("Business name", "business_name", 22, False),
           ("API name", "api_name", 16, False), ("Path segment", "plural", 16, False),
           ("Schema name", "schema_name", 16, False), ("Source object", "source_object", 20, False),
           ("Source container", "source_io", 20, False), ("Operations", "operations", 28, False),
           ("Get one by", "get_by", 10, False), ("Group", "group", 18, False), ("Tags (comma separated)", "tags", 24, False),
           ("Description", "description", 60, False)]
MAP_COLS = [("Namespace", "namespace", 11, True), ("Entity", "entity", 18, True),
            ("Schema name", "schema_name", 14, False), ("JSON path", "json_path", 30, True),
            ("Business name", "business_name", 26, False), ("API name", "api_name", 22, False),
            ("DD attribute reference", "dd_reference", 18, False), ("Data item", "data_item", 24, False),
            ("Type", "type", 10, False), ("Format", "format", 9, False), ("Max length", "max_length", 10, False),
            ("Key", "key", 6, False), ("Required", "required", 9, False), ("Read only", "read_only", 9, False),
            ("System", "system", 8, False), ("Filterable", "filterable", 9, False),
            ("Filter required", "filter_required", 9, False), ("Filter match", "filter_match", 10, False),
            ("PII", "pii", 6, False), ("Money", "money", 7, False), ("Derived", "derived", 8, False),
            ("Enum values", "enum_values", 30, True), ("Source system", "source_system", 16, True),
            ("Source object", "source_object", 16, False),
            ("Source field", "source_field", 26, False), ("Source location", "source_column", 26, False),
            ("Code list / LOV type", "lov_type", 20, False), ("Transformation", "transform", 22, False),
            ("Example", "example", 22, False), ("Description", "description", 44, False),
            ("Column group", "column_group", 18, False), ("Pattern (regex)", "pattern", 26, False)]
REL_COLS = [("From entity", "from_business", 18, True), ("From schema", "from_schema", 14, False),
            ("Cardinality", "cardinality", 13, False), ("To entity", "to_business", 18, True),
            ("To schema", "to_schema", 14, False), ("Business name", "business_name", 22, False),
            ("API name", "api_name", 20, False), ("Exposed as", "expose", 12, False),
            ("Inverse business name", "inverse_business_name", 24, False),
            ("Inverse API name", "inverse_api_name", 24, False), ("Inverse exposed as", "inverse_expose", 16, False),
            ("Join from", "join_from", 30, False), ("Join to", "join_to", 22, False),
            ("Source link", "source_link", 30, False), ("Description", "description", 40, False),
            ("Sub-resource operations", "sub_ops", 22, False), ("Sub-resource key", "sub_key_name", 16, False),
            ("Inverse sub-resource operations", "inverse_sub_ops", 22, False),
            ("Inverse sub-resource key", "inverse_sub_key_name", 16, False),
            ("Set by system", "read_only", 10, False), ("Accepted in requests", "nested_create", 12, False),
            ("Inverse accepted in requests", "inverse_nested_create", 12, False)]
OP_COLS = [("Schema name", "schema_name", 18, False), ("Kind", "kind", 13, False),
           ("Operation", "operation", 16, False), ("Link API name", "link", 18, False),
           ("Method", "method", 8, False), ("Scope", "scope", 10, False), ("Path segment", "path", 16, False),
           ("Summary", "summary", 34, False), ("Request schema", "request", 24, False),
           ("Request form", "request_variant", 11, False), ("Response schema", "response", 22, False),
           ("Response content", "response_content", 22, False),
           ("Query parameters", "query", 34, False),
           ("Status", "status", 7, False), ("If-Match", "if_match", 8, False), ("Errors", "errors", 30, False),
           ("Description", "description", 60, False), ("Operation ID", "operation_id", 26, False)]
HDR_COLS = [("Name", "name", 22, False), ("Direction", "direction", 11, False), ("Required", "required", 9, False),
            ("Type", "type", 9, False), ("Format", "format", 11, False), ("Pattern", "pattern", 16, False),
            ("Example", "example", 30, False), ("Methods", "methods", 18, False), ("Namespaces", "namespaces", 16, False),
            ("Description", "description", 60, False)]
ENUM_COLS = [("Schema name", "schema_name", 14, False), ("Entity", "entity", 18, True),
             ("Column API name", "api_name", 22, False), ("Column", "column", 24, True),
             ("Code", "code", 20, False), ("Label", "label", 22, False), ("Source value", "source_value", 22, False)]
EP_COLS = [("Namespace", "namespace", 11, True), ("Resource", "tag", 18, True), ("Method", "method", 8, True),
           ("Path", "path", 46, True), ("Operation ID", "operation_id", 32, True), ("Summary", "summary", 36, True),
           ("Request schema", "request", 18, True), ("Response schema", "response", 18, True),
           ("Status", "status", 8, True), ("Defined by", "kind", 12, True)]

BOOL_KEYS = {"key", "required", "read_only", "system", "filterable", "filter_required", "pii", "money", "derived"}

# Header names used by 1.x workbooks, still accepted on import.
ALIASES = {"nested create": "nested_create", "tags": "tags", "integration object": "source_io", "source column": "source_column", "lov type": "lov_type"}


def yn(v: bool) -> str:
    return "Y" if v else ""


def parse_bool(v: Any) -> bool:
    return str(v or "").strip().lower() in ("y", "yes", "true", "1", "x")


def json_paths(m: Dict[str, Any], e: Dict[str, Any]) -> List[str]:
    out = []
    if e["operations"]:
        out.append(e["api_name"])
    for owner in m["entities"]:
        for nav in navs_with_target(m, owner):
            if nav["expose"] == "embedded" and nav["target"]["id"] == e["id"]:
                out.append(owner["api_name"] + "." + nav["api_name"] + ("[]" if nav["many"] else ""))
    return out or [e["api_name"]]


def mapping_rows(m: Dict[str, Any]) -> List[Dict[str, Any]]:
    nss = {n["id"]: n for n in m["namespaces"]}
    sss = {x["id"]: x for x in m.get("source_systems", [])}
    rows = []
    for e in m["entities"]:
        n = nss.get(e["namespace_id"], {"key": "?"})
        src = sss.get(e.get("source_system_id") or "")
        prefixes = json_paths(m, e)
        cgs = {g["id"]: g for g in e.get("column_groups") or []}
        for a in e["attributes"]:
            g = cgs.get(a.get("group_id") or "")
            gpath = (g["api_name"] + ("[]" if g.get("many") else "") + ".") if g else ""
            rows.append({
                "namespace": n["key"], "entity": e["business_name"], "schema_name": e["schema_name"],
                "json_path": "\n".join(p + "." + gpath + a["api_name"] for p in prefixes),
                "column_group": (g["api_name"] + ("[]" if g.get("many") else "")) if g else "",
                "pattern": a.get("pattern", ""),
                "business_name": a["business_name"], "api_name": a["api_name"], "type": a["type"],
                "format": a["format"], "max_length": a["max_length"],
                "key": yn(a["key"]), "required": yn(a["required"]), "read_only": yn(a["read_only"]), "system": yn(a.get("system")),
                "dd_reference": a.get("dd_reference", ""), "data_item": a.get("data_item", ""),
                "filterable": yn(a["filterable"]), "pii": yn(a["pii"]), "money": yn(a.get("money")),
                "derived": yn(derived_from(e, a)[0]),
                "filter_required": yn(a.get("filter_required")),
                "filter_match": (a.get("filter_match") or "equals") if a["filterable"] else "",
                "enum_values": ", ".join(v["code"] for v in a["enum"]),
                "source_system": src["name"] if src else "", "source_system_key": src["key"] if src else "",
                "source_object": a.get("source_object") or e["source_object"], "source_field": a["source_field"],
                "source_column": a["source_column"], "lov_type": a["lov_type"], "transform": a["transform"],
                "example": a["example"], "description": a["description"],
            })
    return rows


def _attr_api(e: Optional[Dict[str, Any]], attr_id: str) -> str:
    for a in (e or {}).get("attributes", []):
        if a["id"] == attr_id:
            return a["api_name"]
    return ""


def relationship_rows(m: Dict[str, Any]) -> List[Dict[str, Any]]:
    ents = {e["id"]: e for e in m["entities"]}
    rows = []
    for r in m["relationships"]:
        f, t = ents.get(r["from_entity_id"]), ents.get(r["to_entity_id"])
        rows.append(dict(r, from_business=f["business_name"] if f else "?", from_schema=f["schema_name"] if f else "?",
                         to_business=t["business_name"] if t else "?", to_schema=t["schema_name"] if t else "?",
                         sub_ops=", ".join(r.get("sub_ops") or []),
                         inverse_sub_ops=", ".join(r.get("inverse_sub_ops") or []),
                         sub_key_name=_attr_api(t, r.get("sub_key", "")),
                         inverse_sub_key_name=_attr_api(f, r.get("inverse_sub_key", "")),
                         read_only=yn(r.get("read_only")), nested_create=yn(r.get("nested_create")),
                         inverse_nested_create=yn(r.get("inverse_nested_create"))))
    return rows


def response_content(mode: str, fields: List[str], ent: Optional[Dict[str, Any]], list_property: str = "") -> str:
    """One cell for what a response body holds: blank (default), full, key, none, or column API names;
    for GET lists page, array, or wrapped (optionally 'wrapped: agreements')."""
    if mode == "wrapped" and list_property:
        return "wrapped: " + list_property
    if mode == "fields":
        names = [a["api_name"] for a in (ent or {}).get("attributes", []) if a["id"] in set(fields)]
        return ", ".join(names)
    return mode or ""


def request_form_text(variant: str, fields: List[str], required: List[str], ent: Optional[Dict[str, Any]]) -> str:
    """'fields: reason*, withdrawnBy' for a body of chosen columns (* = required), else the form name."""
    if variant != "fields":
        return variant
    by_id = {a["id"]: a["api_name"] for a in (ent or {}).get("attributes", [])}
    return "fields: " + ", ".join(by_id[f] + ("*" if f in required else "") for f in fields if f in by_id)


def wrapped_property(content: str) -> str:
    m_ = re.match(r"(?i)^wrapped\s*:\s*(\S+)\s*$", content or "")
    return m_.group(1) if m_ else ""


def query_text(specs: Optional[List[Dict[str, Any]]], context: Optional[Dict[str, Any]], m: Dict[str, Any]) -> str:
    """Query parameters as one cell: 'customerNumber (required), status (in), reportedDate (range)'.
    Blank: the columns' filter settings; 'none': no query parameters. A column of another
    entity is written Schema.column; a custom name as 'column as name'."""
    if specs is None:
        return ""
    if not specs:
        return "none"
    where = {a["id"]: (e, a) for e in m["entities"] for a in e["attributes"]}
    parts = []
    for sp in specs:
        hit = where.get(sp.get("attr_id", ""))
        if not hit:
            continue
        e, a = hit
        txt = a["api_name"] if context is not None and e["id"] == context["id"] else e["schema_name"] + "." + a["api_name"]
        if sp.get("name") and sp["name"] != a["api_name"]:
            txt += " as " + sp["name"]
        bits = (["required"] if sp.get("required") else []) + \
               ([sp["match"]] if sp.get("match") and sp["match"] != "equals" else []) + \
               (["header"] if sp.get("location") == "header" else [])
        parts.append(txt + (" (" + ", ".join(bits) + ")" if bits else ""))
    return ", ".join(parts)


def operation_rows(m: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Operation settings, actions and sub-resource operation settings, one row each."""
    ents = {e["id"]: e for e in m["entities"]}

    def sname(eid: str, special: Dict[str, str]) -> str:
        if eid in special:
            return special[eid]
        return ents[eid]["schema_name"] if eid in ents else ""

    rows = []
    for e in m["entities"]:
        for op, ov in (e.get("op_overrides") or {}).items():
            rows.append({"schema_name": e["schema_name"], "kind": "operation", "operation": op,
                         "query": query_text(ov.get("query_params"), e, m),
                         "response_content": response_content(ov.get("response_mode", ""), ov.get("response_fields") or [],
                                                              ents.get(ov.get("response_entity_id", "")) or e,
                                                              ov.get("list_property", "")),
                         "summary": ov.get("summary", ""), "description": ov.get("description", ""),
                         "operation_id": ov.get("operation_id", ""),
                         "request": sname(ov.get("request_entity_id", ""), {"": ""}),
                         "request_variant": request_form_text("fields", ov.get("request_fields") or [], ov.get("request_required") or [],
                                                              ents.get(ov.get("request_entity_id", "")) or e)
                         if ov.get("request_mode") == "fields" else "",
                         "response": sname(ov.get("response_entity_id", ""), {"": "", "none": "none"}),
                         "status": ov.get("status", "")})
        for act in e.get("actions", []):
            rows.append({"schema_name": e["schema_name"], "kind": "action", "operation": act["name"],
                         "query": query_text(act.get("query_params") or [], e, m) if act.get("query_params") else "",
                         "method": act["method"], "scope": act["scope"], "path": act["path"],
                         "summary": act.get("summary", ""), "description": act.get("description", ""),
                         "operation_id": act.get("operation_id", ""),
                         "request": sname(act.get("request_entity_id", ""), {"": ""}),
                         "request_variant": request_form_text(act.get("request_variant", "create"), act.get("request_fields") or [],
                                                              act.get("request_required") or [], ents.get(act.get("request_entity_id", "")))
                         if act.get("request_entity_id") else "",
                         "response": sname(act.get("response_entity_id", ""), {"": "", "self": "self"}),
                         "response_content": response_content(
                             act.get("response_mode", ""), act.get("response_fields") or [],
                             e if act.get("response_entity_id") == "self" else ents.get(act.get("response_entity_id", "")),
                             act.get("list_property", "")),
                         "status": act.get("status", "200"), "if_match": yn(act.get("if_match")),
                         "errors": ", ".join(act.get("errors") or [])})
    for r in m["relationships"]:
        for owner_id, api, key in ((r["from_entity_id"], r["api_name"], "sub_overrides"),
                                   (r["to_entity_id"], r["inverse_api_name"], "inverse_sub_overrides")):
            target = ents.get(r["to_entity_id"] if key == "sub_overrides" else r["from_entity_id"])
            for op, ov in (r.get(key) or {}).items():
                rows.append({"schema_name": sname(owner_id, {}), "kind": "sub-resource", "operation": op, "link": api,
                             "query": query_text(ov.get("query_params"), target, m),
                             "response_content": response_content(ov.get("response_mode", ""), ov.get("response_fields") or [],
                                                                  ents.get(ov.get("response_entity_id", "")) or target,
                                                                  ov.get("list_property", "")),
                             "summary": ov.get("summary", ""), "description": ov.get("description", ""),
                             "operation_id": ov.get("operation_id", ""),
                             "request": sname(ov.get("request_entity_id", ""), {"": ""}),
                             "request_variant": request_form_text("fields", ov.get("request_fields") or [], ov.get("request_required") or [],
                                                                  ents.get(ov.get("request_entity_id", "")) or target)
                             if ov.get("request_mode") == "fields" else "",
                             "response": sname(ov.get("response_entity_id", ""), {"": "", "none": "none"}),
                             "status": ov.get("status", "")})
    return rows


def _sheet(wb: Workbook, title: str, cols, rows: List[Dict[str, Any]], note: str = "") -> Any:
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="5B6B7F", size=9)
        start = 2
    for i, (head, _key, width, calc) in enumerate(cols, 1):
        c = ws.cell(row=start, column=i, value=head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CALC_FILL if calc else HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[start].height = 30
    for r_i, row in enumerate(rows, start + 1):
        for c_i, (_head, key, _w, calc) in enumerate(cols, 1):
            v = row.get(key, "")
            if isinstance(v, list):
                v = ", ".join(v)
            c = ws.cell(row=r_i, column=c_i, value=v if v != "" else None)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
            if calc:
                c.fill = CALC_CELL
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    last = max(start + len(rows), start + 1)
    ws.auto_filter.ref = "A" + str(start) + ":" + get_column_letter(len(cols)) + str(last)
    return ws


def _validation(ws, cols, key: str, options: List[str], first_row: int, rows: int = 500) -> None:
    idx = [c[1] for c in cols].index(key) + 1
    letter = get_column_letter(idx)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "Choose one of: " + ", ".join(options)
    dv.errorTitle = "Not a valid value"
    ws.add_data_validation(dv)
    dv.add(letter + str(first_row) + ":" + letter + str(first_row + rows))


def export_xlsx(m: Dict[str, Any]) -> bytes:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        (m["settings"]["title"], Font(bold=True, size=16, color="1F3A5F")),
        ("API mapping sheet", Font(bold=True, size=12)),
        ("Organisation: " + (m["settings"].get("organisation") or "-"), None),
        ("Source systems: " + (", ".join(x["name"] for x in m.get("source_systems", [])) or "-"), None),
        ("Generated " + dt.datetime.now().strftime("%d %b %Y %H:%M") + " by API Definition Studio " + __version__, None),
        ("", None),
        ("How to use", Font(bold=True)),
        ("• Source systems: the applications data comes from (Siebel, REST services, databases, files, "
         "mainframe, other). Each entity names one by key, or leaves it blank for derived or new data.", None),
        ("• Namespaces: API groupings with their base path, version, security and servers.", None),
        ("• Entities: each source object, its business name, and the name it has in the API "
         "(e.g. business name Contact, API name customer).", None),
        ("• Mapping: one row per column. Business name ↔ API name ↔ DD attribute reference ↔ source field/column. "
         "System = Y: set by the source system, never accepted on POST or PUT (PATCH may still send it "
         "unless Read only is also Y).", None),
        ("• Linked entities: relationships between entities and how each end is exposed "
         "(link, embedded, sub-resource or none).", None),
        ("• Enum values: allowed codes per enum column, with labels and source LOV values.", None),
        ("• Operations: per-operation settings (summary, description, request and response schema, status), "
         "actions such as POST …/{id}/decision, and settings for sub-resource operations. Response schema: blank "
         "means the default; 'none' means no body; for an action, 'self' means the entity itself and blank means "
         "no body. Sub-resource rows name the link by its API name on that entity.", None),
        ("• Headers: standard headers added to every matching operation (request, response, or both). Methods "
         "and Namespaces blank means all. Authorization comes from each namespace's security setting.", None),
        ("• Response content (Operations sheet): blank for the default (POST: "
         + (m.get("settings") or {}).get("post_response", "key") + ", PUT: "
         + (m.get("settings") or {}).get("put_response", "full") + "), full, key (row id only), none, or a "
         "comma-separated list of column API names. GET lists: page ({items, nextCursor} with limit/cursor; default "
         + (m.get("settings") or {}).get("list_response", "page") + "), array (a plain JSON array, no paging), "
         "or wrapped ({agreements: [...]}, no paging; 'wrapped: name' to choose the property).", None),
        ("• Source object (Mapping sheet): the business component a column comes from. It starts as the "
         "entity's; type another one when the column comes from a different business component. Data item: "
         "the data dictionary's data item full name.", None),
        ("• Column group (Mapping sheet): columns sent together under one property: bankDetails puts the column "
         "inside an object, bankDetails[] inside a list of objects. JSON path shows it, e.g. customer.bankDetails.sortCode.", None),
        ("• Pattern (Mapping sheet): a regular expression text values must match, e.g. ^[0-9A-Z]+-[0-9A-Z]+$ for a "
         "Siebel row id. Dates and date-times are given in ISO 8601 (2026-01-31, 2026-01-31T09:30:00Z).", None),
        ("• Derived (Mapping sheet): Y when the column is not held on the entity's own record: it comes from "
         "another business component (always Y then) or is calculated (describe how in Transformation). Shown "
         "highlighted on the data model diagram.", None),
        ("• Money = Y (Mapping sheet): the column is sent as { amount, currency } using the shared Money "
         "schema, e.g. { \"amount\": 2999, \"currency\": \"GBP\" }. The Type is the amount's: integer for pence, "
         "number for decimal pounds. Write the example as 29.99 or 2999 (pence).", None),
        ("• Filters (Mapping sheet): Filterable = Y offers the column as a query parameter on the entity's GET "
         "lists; Filter required = Y makes callers send it; Filter match is equals, in (comma list), range "
         "(From/To) or contains.", None),
        ("• Query parameters (Operations sheet): blank uses the columns' filter settings; 'none' removes them; "
         "otherwise e.g. 'customerNumber (required), status (in), reportedDate (range)'. A column of another entity is "
         "Schema.column; 'customerNumber as clientRef' renames the parameter.", None),
        ("• Get one by (Entities sheet): path for /things/{id}, query for /things?id=.", None),
        ("• Tags (Entities sheet): comma separated, e.g. Agreements, Case. Used to filter the diagram.", None),
        ("• Enum code style (Namespaces sheet): UPPER_SNAKE, kebab-case, camelCase or lower_snake; enum codes in "
         "that namespace are checked against it. Source mapping in OpenAPI = N publishes the contract only "
         "(no x-source-*, DD attribute, data item or transformation).", None),
        ("• Endpoints: the operations the OpenAPI definition contains (generated).", None),
        ("", None),
        ("Dark blue headers can be edited and imported back into the studio. Grey headers are "
         "generated and ignored on import. Y in a flag column means yes; blank means no.", None),
        ("Operations (Entities sheet): any of " + ", ".join(OPERATIONS) + ", comma separated. "
         "Leave blank for an entity that only appears inside another (embedded).", None),
        ("", None),
        ("Source columns by system type (all optional)", Font(bold=True)),
    ]
    lines.append(("Type: Source object | Source container | Source field | Source location | Code list", Font(italic=True)))
    for k, v in KINDS.items():
        lines.append((k + " (" + v["label"] + "): " + " | ".join(v[x] for x in ("object", "container", "field",
                                                                                "location", "code_list")), None))
    for i, (text, font) in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font
    readme.column_dimensions["A"].width = 110

    ss_rows = [dict(x, internal_tokens=", ".join(x.get("internal_tokens") or [])) for x in m.get("source_systems", [])]
    nsk = {n["id"]: n["key"] for n in m["namespaces"]}
    hdr_rows = [dict(hd, required=yn(hd.get("required")), methods=", ".join(hd.get("methods") or []),
                     namespaces=", ".join(nsk.get(x, x) for x in hd.get("namespace_ids") or []))
                for hd in m.get("headers", [])]
    ws = _sheet(wb, "Headers", HDR_COLS, hdr_rows)
    _validation(ws, HDR_COLS, "direction", ["request", "response", "both"], 2)
    _validation(ws, HDR_COLS, "type", ["string", "integer"], 2)
    _validation(ws, HDR_COLS, "format", ["uuid", "date-time", "date"], 2)
    _validation(ws, HDR_COLS, "required", ["Y", "N"], 2)
    ws = _sheet(wb, "Source systems", SS_COLS, ss_rows)
    _validation(ws, SS_COLS, "kind", list(KINDS), 2)

    nss = {n["id"]: n for n in m["namespaces"]}
    sss = {x["id"]: x for x in m.get("source_systems", [])}
    ns_rows = [dict(n, servers="\n".join(s["url"] + (" | " + s["description"] if s["description"] else "")
                                       for s in n["servers"]),
                    enum_style=n.get("enum_style") or "UPPER_SNAKE",
                    source_in_openapi="N" if n.get("source_in_openapi", True) is False else "Y")
               for n in m["namespaces"]]
    ws = _sheet(wb, "Namespaces", NS_COLS, ns_rows)
    _validation(ws, NS_COLS, "security", SECURITY, 2)
    _validation(ws, NS_COLS, "base_path_in", ["paths", "server"], 2)
    _validation(ws, NS_COLS, "enum_style", list(naming.ENUM_STYLES), 2)
    _validation(ws, NS_COLS, "source_in_openapi", ["Y", "N"], 2)

    en_rows = [dict(e, namespace=nss.get(e["namespace_id"], {"key": "?"})["key"],
                    source_system=sss[e["source_system_id"]]["key"] if e.get("source_system_id") in sss else "")
               for e in m["entities"]]
    ws = _sheet(wb, "Entities", EN_COLS, en_rows)
    if sss and len(",".join(x["key"] for x in sss.values())) < 250:
        _validation(ws, EN_COLS, "source_system", [x["key"] for x in sss.values()], 2)

    ws = _sheet(wb, "Mapping", MAP_COLS, mapping_rows(m))
    _validation(ws, MAP_COLS, "type", ATTRIBUTE_TYPES, 2, 2000)
    for k in BOOL_KEYS:
        _validation(ws, MAP_COLS, k, ["Y", "N"], 2, 2000)

    ws = _sheet(wb, "Linked entities", REL_COLS, relationship_rows(m))
    _validation(ws, REL_COLS, "cardinality", CARDINALITIES, 2)
    _validation(ws, REL_COLS, "expose", EXPOSE, 2)
    _validation(ws, REL_COLS, "inverse_expose", EXPOSE, 2)

    enum_rows = []
    for e in m["entities"]:
        for a in e["attributes"]:
            for v in a["enum"]:
                enum_rows.append({"schema_name": e["schema_name"], "entity": e["business_name"],
                                  "api_name": a["api_name"], "column": a["business_name"], **v})
    _sheet(wb, "Enum values", ENUM_COLS, enum_rows)
    ws = _sheet(wb, "Operations", OP_COLS, operation_rows(m))
    _validation(ws, OP_COLS, "kind", ["operation", "action", "sub-resource"], 2)
    _validation(ws, OP_COLS, "method", ["GET", "POST", "PUT", "PATCH", "DELETE"], 2)
    _validation(ws, OP_COLS, "scope", ["item", "collection"], 2)
    _validation(ws, OP_COLS, "status", ["200", "201", "202", "204"], 2)
    _validation(ws, OP_COLS, "if_match", ["Y", "N"], 2)
    _sheet(wb, "Endpoints", EP_COLS, endpoints(m), note="Generated from the model; ignored on import.")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(m: Dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in MAP_COLS])
    for r in mapping_rows(m):
        w.writerow([("" if r.get(k) is None else str(r.get(k)).replace("\n", "; ")) for _h, k, _w, _c in MAP_COLS])
    return buf.getvalue()


# -- import --------------------------------------------------------------------

def _read(ws, cols) -> List[Dict[str, Any]]:
    """Read a sheet by header names (finds the header row in the first 3 rows)."""
    wanted = {h.lower(): k for h, k, _w, calc in cols if not calc}
    keys = set(wanted.values())
    for alias, k in ALIASES.items():
        if k in keys:
            wanted.setdefault(alias, k)
    header_row, index = None, {}
    for r in range(1, 4):
        vals = [str(c.value or "").strip().lower() for c in ws[r]]
        hits = {i: wanted[v] for i, v in enumerate(vals) if v in wanted}
        if len(set(hits.values())) >= max(1, len(keys) // 2):
            header_row, index = r, hits
            break
    if header_row is None:
        raise Invalid("Sheet '" + ws.title + "' has no recognisable header row.")
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rec = {k: row[i] if i < len(row) else None for i, k in index.items()}
        if all(v in (None, "") for v in rec.values()):
            continue
        rows.append({k: ("" if v is None else v) for k, v in rec.items()})
    return rows


def _s(v: Any) -> str:
    if v is None:  # column not in this workbook (e.g. an older version)
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    return str(v).strip()


def import_xlsx(data: bytes, current: Optional[Dict[str, Any]] = None) -> Tuple[Dict[str, Any], List[str]]:
    """Build a model from a mapping workbook. Returns (model, notes)."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface any parse failure to the user
        raise Invalid("Not a readable Excel workbook: " + str(exc))
    need = ["Namespaces", "Entities", "Mapping"]
    missing = [s for s in need if s not in wb.sheetnames]
    if missing:
        raise Invalid("Workbook is missing sheet(s): " + ", ".join(missing))
    notes: List[str] = []
    errors: List[str] = []

    has_systems = "Source systems" in wb.sheetnames
    source_systems = []
    if has_systems:
        for r in _read(wb["Source systems"], SS_COLS):
            key = _s(r.get("key"))
            if not key:
                errors.append("Source systems: a row has no key.")
                continue
            kind = _s(r.get("kind")).lower() or "other"
            if kind not in KINDS:
                errors.append("Source systems: '" + key + "' has unknown type '" + kind + "' (use " + ", ".join(KINDS) + ").")
                continue
            source_systems.append({"id": "ss_" + key.replace("-", "_"), "key": key, "name": _s(r.get("name")), "kind": kind,
                                   "owner": _s(r.get("owner")), "description": _s(r.get("description")),
                                   "internal_tokens": [t.strip() for t in _s(r.get("internal_tokens")).split(",") if t.strip()]})
    else:
        notes.append("No 'Source systems' sheet (a 1.x workbook): every entity is treated as Siebel.")
    ss_by_key = {x["key"]: x for x in source_systems}

    namespaces = []
    for r in _read(wb["Namespaces"], NS_COLS):
        servers = []
        for line in _s(r.get("servers")).splitlines():
            if line.strip():
                url, _, desc = line.partition("|")
                servers.append({"url": url.strip(), "description": desc.strip()})
        namespaces.append({"key": _s(r.get("key")), "title": _s(r.get("title")), "base_path": _s(r.get("base_path")),
                           "base_path_in": _s(r.get("base_path_in")) or "paths",
                           "version": _s(r.get("version")) or "1.0.0", "security": _s(r.get("security")) or "bearer",
                           "owner": _s(r.get("owner")), "description": _s(r.get("description")), "servers": servers,
                           "enum_style": _s(r.get("enum_style")) or "UPPER_SNAKE",
                           "source_in_openapi": _s(r.get("source_in_openapi")).upper() not in ("N", "NO", "FALSE")})
    for n in namespaces:
        n["id"] = "ns_" + n["key"].replace("-", "_")
    ns_by_key = {n["key"]: n for n in namespaces}

    entities: List[Dict[str, Any]] = []
    by_schema: Dict[str, Dict[str, Any]] = {}
    has_tags_col = has_group_col = False
    for r in _read(wb["Entities"], EN_COLS):
        has_tags_col = has_tags_col or "tags" in r
        has_group_col = has_group_col or "group" in r
        nk = _s(r.get("namespace"))
        if nk not in ns_by_key:
            errors.append("Entities: '" + _s(r.get("business_name")) + "' uses unknown namespace '" + nk + "'.")
            continue
        ops = [o.strip().lower() for o in _s(r.get("operations")).replace(";", ",").split(",") if o.strip()]
        sk = _s(r.get("source_system"))
        if sk and sk not in ss_by_key:
            errors.append("Entities: '" + _s(r.get("business_name")) + "' uses unknown source system '" + sk + "'.")
            continue
        e = {"get_by": _s(r.get("get_by")).lower() or "path",
             "namespace_id": ns_by_key[nk]["id"], "source_system_id": ss_by_key[sk]["id"] if sk else "",
             "business_name": _s(r.get("business_name")),
             "api_name": _s(r.get("api_name")), "plural": _s(r.get("plural")), "schema_name": _s(r.get("schema_name")),
             "source_object": _s(r.get("source_object")), "source_io": _s(r.get("source_io")),
             "operations": ops, "description": _s(r.get("description")), "attributes": [],
             "tags": clean_tags(_s(r.get("tags"))), "group": _s(r.get("group"))}
        if not e["business_name"]:
            errors.append("Entities: a row has no business name.")
            continue
        entities.append(e)

    # derive names now so Mapping rows can refer to schema names that were left blank
    from .store import fill_defaults
    tmp = fill_defaults({"namespaces": namespaces, "entities": entities, "relationships": []})
    entities = tmp["entities"]
    for e in entities:
        by_schema[e["schema_name"].lower()] = e
        by_schema.setdefault(e["business_name"].lower(), e)

    attr_index: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for i, r in enumerate(_read(wb["Mapping"], MAP_COLS), 1):
        sch = _s(r.get("schema_name")).lower()
        e = by_schema.get(sch)
        if not e:
            errors.append("Mapping row " + str(i) + ": unknown schema '" + _s(r.get("schema_name")) + "'.")
            continue
        ml = _s(r.get("max_length"))
        a = {"business_name": _s(r.get("business_name")), "api_name": _s(r.get("api_name")),
             "type": _s(r.get("type")).lower() or "string", "format": _s(r.get("format")),
             "max_length": int(float(ml)) if ml else None,
             "source_field": _s(r.get("source_field")), "source_column": _s(r.get("source_column")),
             "lov_type": _s(r.get("lov_type")), "transform": _s(r.get("transform")),
             "filter_match": _s(r.get("filter_match")).lower() or "equals",
             "dd_reference": _s(r.get("dd_reference")), "data_item": _s(r.get("data_item")),
             "pattern": str(r.get("pattern") or "").strip(),
             "example": _s(r.get("example")), "description": _s(r.get("description")), "enum": []}
        # column group: bankDetails (an object) or bankDetails[] (a list); kept by name, made if new
        cg = _s(r.get("column_group")).replace(" ", "")
        if cg:
            many = cg.endswith("[]")
            gname = cg[:-2] if many else cg
            groups_ = e.setdefault("column_groups", [])
            g = next((x for x in groups_ if x["api_name"] == gname), None)
            if not g:
                old = next((x for ce in (current or {}).get("entities", []) if ce["schema_name"].lower() == e["schema_name"].lower()
                            for x in ce.get("column_groups") or [] if x["api_name"] == gname), None)
                g = dict(old or {"business_name": "", "required": False, "description": ""}, api_name=gname, many=many)
                g["id"] = g.get("id") or new_id("cg")
                groups_.append(g)
            a["group_id"] = g["id"]
        col_bc = _s(r.get("source_object"))   # the column's business component; only kept when not the entity's
        a["source_object"] = col_bc if col_bc and col_bc.lower() != (e.get("source_object") or "").lower() else ""
        for k in BOOL_KEYS:
            a[k] = parse_bool(r.get(k))
        if a["derived"] and a["source_object"]:
            a["derived"] = False   # another business component already makes it derived; keep the flag for calculated ones
        if a["type"] == "money":   # a 1.14 workbook: money was a type
            a["type"], a["money"] = "integer", True
        if not a["business_name"] and not a["api_name"]:
            errors.append("Mapping row " + str(i) + ": no business name or API name.")
            continue
        e["attributes"].append(a)

    tmp = fill_defaults({"namespaces": namespaces, "entities": entities, "relationships": []})
    entities = tmp["entities"]
    by_schema = {}
    for e in entities:
        by_schema[e["schema_name"].lower()] = e
        by_schema.setdefault(e["business_name"].lower(), e)
        for a in e["attributes"]:
            attr_index[(e["schema_name"].lower(), a["api_name"])] = a

    if "Enum values" in wb.sheetnames:
        for i, r in enumerate(_read(wb["Enum values"], ENUM_COLS), 1):
            a = attr_index.get((_s(r.get("schema_name")).lower(), _s(r.get("api_name"))))
            if not a:
                errors.append("Enum values row " + str(i) + ": no column " + _s(r.get("schema_name")) + "."
                              + _s(r.get("api_name")) + ".")
                continue
            a["enum"].append({"code": _s(r.get("code")), "label": _s(r.get("label")),
                              "source_value": _s(r.get("source_value"))})
        for (sch, api), a in attr_index.items():
            if a["enum"] and a["type"] != "enum":
                notes.append(sch + "." + api + " has enum values but type " + a["type"] + "; values ignored.")

    relationships = []
    if "Linked entities" in wb.sheetnames:
        for i, r in enumerate(_read(wb["Linked entities"], REL_COLS), 1):
            f = by_schema.get(_s(r.get("from_schema")).lower())
            t = by_schema.get(_s(r.get("to_schema")).lower())
            if not f or not t:
                errors.append("Linked entities row " + str(i) + ": unknown entity '"
                              + (_s(r.get("from_schema")) if not f else _s(r.get("to_schema"))) + "'.")
                continue
            def key_id(ent: Dict[str, Any], api: str) -> str:
                return next((a["id"] for a in ent["attributes"] if a["api_name"] == api), "") if api else ""

            def ops_list(v: Any) -> List[str]:
                return [x.strip().lower() for x in _s(v).replace(";", ",").split(",") if x.strip()]
            sk, isk = _s(r.get("sub_key_name")), _s(r.get("inverse_sub_key_name"))
            if sk and not key_id(t, sk):
                errors.append("Linked entities row " + str(i) + ": sub-resource key '" + sk + "' is not a column of "
                              + t["schema_name"] + ".")
            if isk and not key_id(f, isk):
                errors.append("Linked entities row " + str(i) + ": inverse sub-resource key '" + isk
                              + "' is not a column of " + f["schema_name"] + ".")
            relationships.append({
                "sub_ops": ops_list(r.get("sub_ops")), "sub_key": key_id(t, sk),
                "inverse_sub_ops": ops_list(r.get("inverse_sub_ops")), "inverse_sub_key": key_id(f, isk),
                "read_only": parse_bool(r.get("read_only")), "nested_create": parse_bool(r.get("nested_create")),
                "inverse_nested_create": parse_bool(r.get("inverse_nested_create")),
                "sub_overrides": {}, "inverse_sub_overrides": {},
                "from_entity_id": f["id"], "to_entity_id": t["id"],
                "cardinality": _s(r.get("cardinality")) or "many-to-one",
                "business_name": _s(r.get("business_name")), "api_name": _s(r.get("api_name")),
                "expose": _s(r.get("expose")) or "none",
                "inverse_business_name": _s(r.get("inverse_business_name")),
                "inverse_api_name": _s(r.get("inverse_api_name")),
                "inverse_expose": _s(r.get("inverse_expose")) or "none",
                "join_from": _s(r.get("join_from")), "join_to": _s(r.get("join_to")),
                "source_link": _s(r.get("source_link")), "description": _s(r.get("description"))})
    else:
        notes.append("No 'Linked entities' sheet; the model will have no links.")

    if "Operations" in wb.sheetnames:
        def ent_of(name: str) -> Optional[Dict[str, Any]]:
            return by_schema.get(name.lower()) if name else None

        for i, r in enumerate(_read(wb["Operations"], OP_COLS), 1):
            where = "Operations row " + str(i)
            owner = ent_of(_s(r.get("schema_name")))
            kind = _s(r.get("kind")).lower()
            op = _s(r.get("operation"))
            if not owner:
                errors.append(where + ": unknown schema '" + _s(r.get("schema_name")) + "'.")
                continue

            def ref_id(v: str, special: Tuple[str, ...]) -> Optional[str]:
                if v in special:
                    return v
                e2 = ent_of(v)
                if not e2:
                    errors.append(where + ": unknown schema '" + v + "'.")
                    return None
                return e2["id"]
            req = ref_id(_s(r.get("request")), ("",))
            content = _s(r.get("response_content"))
            qtext = _s(r.get("query"))

            def parse_query(context: Optional[Dict[str, Any]]) -> Optional[List[Dict[str, Any]]]:
                if not qtext:
                    return None
                if qtext.lower() == "none":
                    return []
                specs = []
                for part in re.findall(r"[^,(]+(?:\([^)]*\))?", qtext):
                    mm = re.match(r"^\s*([A-Za-z_][\w.]*)(?:\s+as\s+([\w-]+))?\s*(?:\(([^)]*)\))?\s*$", part)
                    if not part.strip():
                        continue
                    if not mm:
                        errors.append(where + ": could not read query parameter '" + part.strip() + "'.")
                        continue
                    col, alias, opts = mm.group(1), mm.group(2) or "", mm.group(3) or ""
                    ent = context
                    if "." in col:
                        sch_name, col = col.split(".", 1)
                        ent = ent_of(sch_name)
                    a = next((x for x in (ent or {}).get("attributes", []) if x["api_name"] == col), None)
                    if not a:
                        errors.append(where + ": query parameter column '" + mm.group(1) + "' not found.")
                        continue
                    words_ = [w.strip().lower() for w in opts.split(",") if w.strip()]
                    match = next((w for w in words_ if w in ("equals", "in", "range", "contains")), "equals")
                    specs.append({"attr_id": a["id"], "required": "required" in words_, "match": match, "name": alias,
                                  "location": "header" if "header" in words_ else "query"})
                return specs

            def parse_request_form(ent: Optional[Dict[str, Any]]) -> Tuple[str, List[str], List[str]]:
                """'fields: reason*, withdrawnBy' -> (fields, ids, required ids); else (form, [], [])."""
                txt = _s(r.get("request_variant"))
                mm = re.match(r"(?i)^fields\s*:\s*(.*)$", txt)
                if not mm:
                    return txt.lower(), [], []
                by_api = {a["api_name"]: a["id"] for a in (ent or {}).get("attributes", [])}
                ids, req_ids = [], []
                for part in [x.strip() for x in mm.group(1).split(",") if x.strip()]:
                    nm_ = part.rstrip("*").strip()
                    if nm_ not in by_api:
                        errors.append(where + ": request field '" + nm_ + "' is not a column of "
                                      + ((ent or {}).get("schema_name") or "the request entity") + ".")
                        continue
                    ids.append(by_api[nm_])
                    if part.endswith("*"):
                        req_ids.append(by_api[nm_])
                return "fields", ids, req_ids

            def mode_fields(ent: Optional[Dict[str, Any]]) -> Tuple[str, List[str]]:
                if content.lower() in ("", "full", "key", "none", "page", "array", "wrapped"):
                    return content.lower(), []
                if re.match(r"(?i)^wrapped\s*:", content):
                    return "wrapped", []
                names = [x.strip() for x in content.split(",") if x.strip()]
                by_api = {a["api_name"]: a["id"] for a in (ent or {}).get("attributes", [])}
                missing = [x for x in names if x not in by_api]
                if missing:
                    errors.append(where + ": response content names columns not in "
                                  + ((ent or {}).get("schema_name") or "the response") + ": " + ", ".join(missing) + ".")
                return "fields", [by_api[x] for x in names if x in by_api]
            if kind in ("operation", "sub-resource"):
                resp = ref_id(_s(r.get("response")), ("", "none"))
                ov = {"status": _s(r.get("status")), "request_entity_id": req or "", "response_entity_id": resp or "",
                      "summary": _s(r.get("summary")), "description": _s(r.get("description")),
                      "operation_id": _s(r.get("operation_id"))}
                def apply_request_form(default_ent: Optional[Dict[str, Any]]) -> None:
                    rv_, rf_, rr_ = parse_request_form(next((x for x in entities if x["id"] == req), None) or default_ent)
                    if rv_ == "fields":
                        ov["request_mode"], ov["request_fields"], ov["request_required"] = "fields", rf_, rr_
                resp_ent = next((x for x in entities if x["id"] == resp), None) if resp not in ("", "none", None) else None
                if kind == "operation":
                    apply_request_form(owner)
                    ov["response_mode"], ov["response_fields"] = mode_fields(resp_ent or owner)
                    ov["list_property"] = wrapped_property(content)
                    ov["query_params"] = parse_query(owner)
                    owner.setdefault("op_overrides", {})[op.lower()] = ov
                    continue
                link = _s(r.get("link"))
                hit = None
                for rel in relationships:
                    if rel["from_entity_id"] == owner["id"] and rel["api_name"] == link:
                        hit = (rel, "sub_overrides")
                    elif rel["to_entity_id"] == owner["id"] and rel["inverse_api_name"] == link:
                        hit = (rel, "inverse_sub_overrides")
                    if hit:
                        break
                if not hit:
                    errors.append(where + ": " + owner["schema_name"] + " has no link named '" + link + "'.")
                    continue
                tgt_id = hit[0]["to_entity_id"] if hit[1] == "sub_overrides" else hit[0]["from_entity_id"]
                ov["response_mode"], ov["response_fields"] = mode_fields(
                    resp_ent or next((x for x in entities if x["id"] == tgt_id), None))
                ov["list_property"] = wrapped_property(content)
                ov["query_params"] = parse_query(next((x for x in entities if x["id"] == tgt_id), None))
                apply_request_form(next((x for x in entities if x["id"] == tgt_id), None))
                hit[0][hit[1]][op.lower()] = ov
            elif kind == "action":
                resp = ref_id(_s(r.get("response")), ("", "self"))
                a_mode, a_fields = mode_fields(owner if resp == "self" else next((x for x in entities if x["id"] == resp), None))
                rv, rf, rr = parse_request_form(next((x for x in entities if x["id"] == req), None))
                owner.setdefault("actions", []).append({
                    "request_fields": rf, "request_required": rr,
                    "name": op, "method": _s(r.get("method")).upper() or "POST",
                    "scope": _s(r.get("scope")).lower() or "item", "path": _s(r.get("path")),
                    "summary": _s(r.get("summary")), "description": _s(r.get("description")),
                    "operation_id": _s(r.get("operation_id")),
                    "request_entity_id": req or "", "request_variant": rv or "create",
                    "response_entity_id": resp if resp is not None else "self", "status": _s(r.get("status")) or "200",
                    "response_mode": a_mode, "response_fields": a_fields, "list_property": wrapped_property(content), "query_params": parse_query(owner) or [],
                    "if_match": parse_bool(r.get("if_match")),
                    "errors": [x.strip() for x in _s(r.get("errors")).split(",") if x.strip()]})
            else:
                errors.append(where + ": kind must be operation, action or sub-resource.")

    if errors:
        raise Invalid("; ".join(errors[:20]) + (" (and " + str(len(errors) - 20) + " more)" if len(errors) > 20 else ""))
    headers = (current or {}).get("headers", [])
    if "Headers" in wb.sheetnames:
        ns_by_key2 = {n["key"]: n["id"] for n in namespaces}
        headers = []
        for i, r in enumerate(_read(wb["Headers"], HDR_COLS), 1):
            nm = _s(r.get("name"))
            if not nm:
                continue
            nks = [x.strip() for x in _s(r.get("namespaces")).split(",") if x.strip()]
            unknown = [x for x in nks if x not in ns_by_key2]
            if unknown:
                errors.append("Headers row " + str(i) + ": unknown namespace " + ", ".join(unknown) + ".")
            headers.append({"name": nm, "direction": _s(r.get("direction")).lower() or "request",
                            "required": parse_bool(r.get("required")), "type": _s(r.get("type")).lower() or "string",
                            "format": _s(r.get("format")).lower(), "pattern": _s(r.get("pattern")),
                            "example": _s(r.get("example")), "description": _s(r.get("description")),
                            "methods": [x.strip().upper() for x in _s(r.get("methods")).split(",") if x.strip()],
                            "namespace_ids": [ns_by_key2[x] for x in nks if x in ns_by_key2]})
    elif current is not None:
        notes.append("No 'Headers' sheet: kept the current model's headers.")
    # a workbook from before 1.7 has no Tags column: keep the current tags
    cur_tags = {e["schema_name"].lower(): e.get("tags") or [] for e in (current or {}).get("entities", [])}
    kept = 0
    for e in entities:
        if not has_tags_col and cur_tags.get(e["schema_name"].lower()):
            e["tags"] = list(cur_tags[e["schema_name"].lower()])
            kept += 1
    cur_groups = {e["schema_name"].lower(): e.get("group") or "" for e in (current or {}).get("entities", [])}
    if not has_group_col:
        for e in entities:
            e["group"] = cur_groups.get(e["schema_name"].lower(), "")
    if kept:
        notes.append("No Tags column: kept the current tags on " + str(kept) + " entit" + ("y." if kept == 1 else "ies."))
    model = {"schema_version": (3 if has_systems else 1), "headers": headers, "source_systems": source_systems,
             "settings": (current or {}).get("settings") or {"title": "API catalogue"},
             "namespaces": namespaces, "entities": entities, "relationships": relationships}
    return fill_defaults(model), notes
